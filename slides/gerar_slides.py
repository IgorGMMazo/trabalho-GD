# ================================================================
#  GÊMEO DIGITAL — GERADOR DOS SLIDES (.pptx)
#
#  Gera slides/gemeo_digital_granja.pptx cobrindo todos os tópicos
#  exigidos pela disciplina, com ênfase em ANÁLISE CRÍTICA.
#
#  Requisitos:  pip install python-pptx
#  Uso:         python slides/gerar_slides.py
# ================================================================

from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

# ── Paleta (mesma do dashboard) ─────────────────────────────────
BG      = RGBColor(0x0B, 0x12, 0x20)
PANEL   = RGBColor(0x12, 0x1B, 0x2E)
ACCENT  = RGBColor(0x38, 0xBD, 0xF8)
TXT     = RGBColor(0xE6, 0xED, 0xF7)
MUTED   = RGBColor(0x8A, 0xA0, 0xC0)
VERDE   = RGBColor(0x22, 0xC5, 0x5E)
AMAREL  = RGBColor(0xEA, 0xB3, 0x08)
VERMEL  = RGBColor(0xEF, 0x44, 0x44)
LARANJA = RGBColor(0xF9, 0x73, 0x16)

EMU_W, EMU_H = Emu(12192000), Emu(6858000)   # 16:9


def _fundo(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def _caixa(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def _barra(slide):
    # faixa de destaque no topo
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), EMU_W, Emu(120000))
    s.fill.solid(); s.fill.fore_color.rgb = ACCENT
    s.line.fill.background()
    return s


def slide_titulo(prs, titulo, subtitulo, autores):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fundo(s)
    from pptx.enum.shapes import MSO_SHAPE
    faixa = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(2400000), EMU_W, Emu(28000))
    faixa.fill.solid(); faixa.fill.fore_color.rgb = ACCENT; faixa.line.fill.background()

    _, tf = _caixa(s, 700000, 2650000, 10800000, 2000000)
    p = tf.paragraphs[0]; p.text = titulo
    p.font.size = Pt(40); p.font.bold = True; p.font.color.rgb = TXT
    p2 = tf.add_paragraph(); p2.text = subtitulo
    p2.font.size = Pt(20); p2.font.color.rgb = ACCENT

    _, tf2 = _caixa(s, 700000, 700000, 10800000, 700000)
    p = tf2.paragraphs[0]; p.text = "🐔  GÊMEO DIGITAL"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = MUTED

    _, tf3 = _caixa(s, 700000, 5200000, 10800000, 1200000)
    p = tf3.paragraphs[0]; p.text = autores
    p.font.size = Pt(15); p.font.color.rgb = MUTED


def slide_secao(prs, numero, titulo):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fundo(s)
    _, tf = _caixa(s, 700000, 2700000, 10800000, 1600000)
    p = tf.paragraphs[0]; p.text = f"{numero:02d}"
    p.font.size = Pt(54); p.font.bold = True; p.font.color.rgb = ACCENT
    p2 = tf.add_paragraph(); p2.text = titulo
    p2.font.size = Pt(30); p2.font.bold = True; p2.font.color.rgb = TXT


def slide_conteudo(prs, titulo, bullets, rodape=None):
    """bullets: lista de (texto, nivel, cor_opcional)."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fundo(s); _barra(s)

    _, tf = _caixa(s, 600000, 320000, 11000000, 900000)
    p = tf.paragraphs[0]; p.text = titulo
    p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = TXT

    _, body = _caixa(s, 650000, 1350000, 10900000, 4900000)
    body.word_wrap = True
    first = True
    for item in bullets:
        texto, nivel = item[0], item[1]
        cor = item[2] if len(item) > 2 else (TXT if nivel == 0 else MUTED)
        p = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        marca = "" if nivel < 0 else ("●  " if nivel == 0 else "–  ")
        p.text = marca + texto
        p.level = max(0, nivel)
        p.font.size = Pt(20 if nivel == 0 else 16) if nivel >= 0 else Pt(18)
        p.font.bold = (nivel == 0)
        p.font.color.rgb = cor
        p.space_after = Pt(7)

    if rodape:
        _, rf = _caixa(s, 600000, 6350000, 11000000, 380000)
        p = rf.paragraphs[0]; p.text = rodape
        p.font.size = Pt(11); p.font.italic = True; p.font.color.rgb = MUTED
    return s


def slide_mono(prs, titulo, linhas, rodape=None):
    """Slide com bloco monoespaçado (diagramas/arquitetura)."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fundo(s); _barra(s)
    _, tf = _caixa(s, 600000, 320000, 11000000, 900000)
    p = tf.paragraphs[0]; p.text = titulo
    p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = TXT

    from pptx.enum.shapes import MSO_SHAPE
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(650000), Emu(1350000),
                              Emu(10900000), Emu(4700000))
    card.fill.solid(); card.fill.fore_color.rgb = PANEL; card.line.color.rgb = ACCENT
    card.line.width = Pt(1)
    tf2 = card.text_frame; tf2.word_wrap = True
    tf2.vertical_anchor = MSO_ANCHOR.TOP
    tf2.margin_left = Emu(250000); tf2.margin_top = Emu(200000)
    first = True
    for ln in linhas:
        p = tf2.paragraphs[0] if first else tf2.add_paragraph()
        first = False
        p.text = ln
        p.font.size = Pt(15); p.font.name = "Consolas"; p.font.color.rgb = TXT
    if rodape:
        _, rf = _caixa(s, 600000, 6300000, 11000000, 380000)
        p = rf.paragraphs[0]; p.text = rodape
        p.font.size = Pt(11); p.font.italic = True; p.font.color.rgb = MUTED
    return s


# ════════════════════════════════════════════════════════════════
#  CONTEÚDO DOS SLIDES
# ════════════════════════════════════════════════════════════════
def construir(prs):
    AUTORES = "Disciplina de Gêmeos Digitais — UFG   |   Equipe: Daniel Rios, Igor G. M. Mazo   |   2026"

    # 0 — Capa
    slide_titulo(prs,
        "Gêmeo Digital de uma Granja Avícola",
        "Coordenação de múltiplos atuadores em um ambiente físico compartilhado, guiada por bem-estar animal",
        AUTORES)

    # 1 — Contextualização
    slide_conteudo(prs, "1 · Contextualização do problema", [
        ("Na avicultura intensiva, o AMBIENTE do galpão determina diretamente bem-estar, sanidade e produtividade das aves.", 0),
        ("Três variáveis ambientais são críticas e precisam ser controladas ao mesmo tempo:", 0),
        ("LUZ (lux) — regula consumo de ração, descanso e comportamento.", 1),
        ("CLIMA (temperatura + umidade → ITU) — conforto térmico; pintinhos exigem ~32 °C.", 1),
        ("GASES (amônia, NH₃) — acima de 25 ppm causa lesão respiratória e mortalidade.", 1),
        ("O problema central: esses eixos NÃO são independentes — os atuadores compartilham o mesmo ar.", 0, ACCENT),
        ("Acionar o exaustor para limpar amônia RESFRIA o galpão e pode matar pintinhos de frio.", 1, VERMEL),
    ], rodape="Um controle ingênuo (três malhas isoladas) pode 'consertar' um eixo e destruir outro.")

    # 2 — Pergunta de pesquisa
    slide_conteudo(prs, "2 · Pergunta de pesquisa", [
        ("Como um gêmeo digital REATIVO pode coordenar múltiplos atuadores que", -1, TXT),
        ("compartilham o mesmo ambiente físico de uma granja — resolvendo os", -1, ACCENT),
        ("conflitos entre os eixos luz, clima e gases — SEM comprometer o", -1, ACCENT),
        ("bem-estar animal?", -1, ACCENT),
        ("", -1),
        ("Sub-perguntas:", 0),
        ("Como representar os efeitos cruzados de um atuador sobre eixos que não são o seu?", 1),
        ("Que critério usar para arbitrar quando dois eixos pedem ações opostas?", 1),
    ])

    # 3 — Objetivos
    slide_conteudo(prs, "3 · Objetivos do trabalho", [
        ("Objetivo geral:", 0),
        ("Construir um gêmeo digital reativo da granja, com os três eixos integrados sobre um ambiente físico compartilhado e um supervisor que arbitra conflitos por bem-estar animal.", 1),
        ("Objetivos específicos:", 0),
        ("Replicar fielmente, em Python, o firmware dos sensores/atuadores (ESP32/Wokwi).", 1),
        ("Modelar a física do galpão capturando os acoplamentos entre atuadores.", 1),
        ("Implementar a hierarquia de decisão: térmico letal > amônia tóxica > luz.", 1),
        ("Visualizar tudo em um dashboard de gêmeo digital (certo/errado + mitigação).", 1),
        ("Validar por cenários cíclicos e por uma suíte de testes automatizados.", 1),
    ])

    # 4 — Motivação e relevância
    slide_conteudo(prs, "4 · Motivação e relevância da proposta", [
        ("A avicultura é central na produção de proteína; pequenas falhas ambientais geram grandes perdas.", 0),
        ("Amônia e estresse térmico estão entre as principais causas de queda de desempenho e mortalidade.", 1),
        ("Sistemas comerciais reais já têm exaustores, aquecedores, cortinas e iluminação — mas frequentemente controlados de forma isolada.", 0),
        ("Um gêmeo digital permite RACIOCINAR sobre o efeito sistêmico antes de atuar:", 0, ACCENT),
        ("testar 'e se?' sem risco às aves;", 1),
        ("explicitar conflitos que um operador humano resolveria por intuição;", 1),
        ("servir de base para evolução preditiva (manejo proativo).", 1),
    ])

    # 5 — Fundamentação teórica
    slide_conteudo(prs, "5 · Fundamentação teórica e trabalhos relacionados", [
        ("Gêmeo digital: réplica virtual de um sistema físico, alimentada por dados, usada para monitorar e decidir.", 0),
        ("Bases zootécnicas adotadas (limiares):", 0),
        ("NH₃: < 10 ppm saudável | 20–25 limite regulatório | > 25 lesão respiratória.", 1),
        ("ITU (índice de Thom): ITU = T + 0,36·Tpo + 41,5; < 74 conforto, > 78 crítico.", 1),
        ("Curva de aquecimento do pinto: ~32 °C na chegada, reduzindo até ~21 °C.", 1),
        ("Iluminância: alta na chegada (estímulo ao consumo), baixa no crescimento.", 1),
        ("Controle: controlador proporcional (P) com banda morta; IoT/MQTT em aviários.", 0),
        ("Trabalhos relacionados: automação de aviários climatizados, sensoriamento IoT e índices de conforto térmico animal.", 0),
    ], rodape="Refs.: Orffa (2023); Precision Poultry Farming/UGA; Frontiers Vet. Sci. (2020); curva de aquecimento Cobb/Ross.")

    # 6 — Metodologia
    slide_conteudo(prs, "6 · Metodologia adotada", [
        ("Abordagem incremental, partindo do que já existia:", 0),
        ("Eixo GASES (branch main) — MQ-135 → exaustor, via MQTT.", 1),
        ("Eixo LUZ (branch main_v2) — BH1750 → lâmpada, arquitetura modular (Bus/Controller).", 1),
        ("Eixo CLIMA (novos projetos Wokwi) — DHT22 → ITU → aquecedor/cortina/alerta.", 1),
        ("Passos:", 0),
        ("1) Replicar os firmwares em Python (fórmulas e ruído idênticos).", 1),
        ("2) Unificar os três eixos sobre UM ambiente físico compartilhado.", 1),
        ("3) Inserir o supervisor de conflitos guiado por bem-estar.", 1),
        ("4) Validar por cenários cíclicos + testes automatizados (pytest).", 1),
    ])

    # 7 — Arquitetura
    slide_mono(prs, "7 · Arquitetura da solução", [
        "        ┌─────────────────── AMBIENTE FÍSICO COMPARTILHADO ───────────────────┐",
        "        │   temperatura · umidade · NH₃ · luz natural   (modelo dinâmico)      │",
        "        └──────▲───────────────────────────────────────────────────▲──────────┘",
        "               │ lê                                       aplica   │",
        "        ┌──────┴───────┐     ┌───────────────┐     ┌───────────────┴─────────┐",
        "        │  SENSORES    │ ──► │ CONTROLADORES │ ──► │      SUPERVISOR          │",
        "        │ lux·clima·gas│     │ 1 por eixo (P)│     │ arbitra conflitos por    │",
        "        └──────────────┘     └───────────────┘     │ BEM-ESTAR e comanda os   │",
        "                                                   │ atuadores compartilhados │",
        "                                                   └──────────────────────────┘",
        "                                                            │",
        "      ATUADORES:  🌀 exaustor (gases+clima)  🔥 aquecedor  💡 lâmpada  🪟 cortina",
        "",
        "   ciclo fechado:  ambiente → sensores → controladores → supervisor → atuadores → ambiente",
    ], rodape="O supervisor é o que diferencia o gêmeo de três malhas isoladas: ele enxerga o sistema inteiro.")

    # 8 — Efeitos cruzados (matriz)
    slide_conteudo(prs, "8 · O núcleo do problema: efeitos cruzados dos atuadores", [
        ("Exaustor (resolve GASES) → ↓ temperatura e ↓ umidade — pode gelar pintinhos.", 0, VERMEL),
        ("Aquecedor (resolve FRIO) → ↑ volatilização de NH₃ da cama — piora os GASES.", 0, LARANJA),
        ("Cortina/ventilação (resolve CALOR) → mesmo efeito do exaustor: resfria e limpa ar.", 0, AMAREL),
        ("Lâmpada (resolve LUZ) → adiciona calor ao galpão (efeito menor, relevante em incandescente).", 0, MUTED),
        ("", -1),
        ("Consequência: o EXAUSTOR é um recurso ÚNICO disputado por gases e clima;", 0, ACCENT),
        ("e aquecedor × exaustor são fisicamente OPOSTOS. Alguém precisa arbitrar.", 0, ACCENT),
    ])

    # 9 — Supervisor / hierarquia
    slide_conteudo(prs, "9 · Supervisor: hierarquia de bem-estar animal", [
        ("Quando há conflito, a decisão segue a gravidade do risco às aves:", 0),
        ("1º  Risco térmico letal — frio em pintinhos / calor extremo.", 1, VERMEL),
        ("2º  Amônia tóxica — NH₃ ≥ 25 ppm.", 1, LARANJA),
        ("3º  Conforto lumínico — luz.", 1, AMAREL),
        ("Regra de ouro:", 0, ACCENT),
        ("NUNCA resfriar além do 'teto seguro' da fase para combater amônia.", 1),
        ("Em vez disso: ventila o tolerável e COMPENSA com aquecedor, registrando o conflito.", 1),
        ("Sinergia: no calor, ventilar resfria E remove NH₃ — um atuador resolve dois eixos.", 0, VERDE),
    ])

    # 10 — Hardware/software
    slide_conteudo(prs, "10 · Hardware, software, ferramentas e tecnologias", [
        ("Hardware (referência, simulado no Wokwi):", 0),
        ("ESP32 DevKit; BH1750 (lux); DHT22 (temp/umid); MQ-135 (NH₃); LEDs; servo (cortina).", 1),
        ("Software do gêmeo digital:", 0),
        ("Python 3 (apenas biblioteca padrão) — sem dependências para rodar.", 1),
        ("Servidor HTTP + Server-Sent Events (SSE) embutido; dashboard em HTML/CSS/Canvas/SVG.", 1),
        ("Ferramentas:", 0),
        ("Wokwi (firmware/diagramas); MQTT/HiveMQ (opcional, p/ hardware real); pytest (testes); Git/VS Code.", 1),
        ("Decisão de projeto: rodar 100% offline (sem internet, sem broker) — confiabilidade na banca.", 0, ACCENT),
    ])

    # 11 — Etapas de desenvolvimento
    slide_conteudo(prs, "11 · Etapas de desenvolvimento e implementação", [
        ("Levantamento e leitura dos três firmwares (gases, luz, clima).", 0),
        ("Réplica em Python: fórmulas de ITU/ponto de orvalho e ruído gaussiano idênticos ao sketch.", 0),
        ("Modelo físico do galpão (EDOs de 1ª ordem) com os acoplamentos entre atuadores.", 0),
        ("Controladores P por eixo + supervisor de arbitragem.", 0),
        ("Motor de cenários cíclicos + dashboard de gêmeo digital (SSE).", 0),
        ("Validação: 22 testes automatizados (física, controladores e, sobretudo, conflitos).", 0),
        ("Correções feitas durante o caminho (ver 'Análise crítica').", 0, ACCENT),
    ])

    # 12 — Base de dados
    slide_conteudo(prs, "12 · Base de dados utilizada", [
        ("Este projeto NÃO utiliza banco de dados.", 0, ACCENT),
        ("Justificativa de projeto: o gêmeo é REATIVO ao estado instantâneo do ambiente —", 0),
        ("decide a cada ciclo a partir das leituras atuais, sem necessidade de histórico persistido.", 1),
        ("Mantemos apenas uma janela curta em memória (para os gráficos ao vivo).", 1),
        ("Persistir séries temporais e treinar modelos preditivos é, conscientemente,", 0),
        ("um TRABALHO FUTURO (ver slide de Trabalhos Futuros).", 1, AMAREL),
    ])

    # 13 — Experimentos
    slide_conteudo(prs, "13 · Experimentos realizados (cenários cíclicos)", [
        ("O motor de cenários percorre situações que estressam cada eixo e os conflitos:", 0),
        ("Operação normal — clima ameno, cama limpa.", 1, VERDE),
        ("Acúmulo de amônia — cama saturada eleva NH₃; exaustor reage.", 1, AMAREL),
        ("Onda de calor — ITU sobe; ventilação máxima + sinergia com gases.", 1, LARANJA),
        ("Noite fria — temperatura cai; aquecedor protege as aves.", 1, ACCENT),
        ("Conflito crítico: frio + amônia tóxica — o supervisor arbitra para não matar de frio.", 1, VERMEL),
        ("Repetidos para as 3 fases de vida (chegada, crescimento, final).", 0),
    ])

    # 14 — Métricas
    slide_conteudo(prs, "14 · Métricas de avaliação adotadas", [
        ("Como o sistema é reativo (sem ML), as métricas medem QUALIDADE DA DECISÃO:", 0),
        ("Violações do teto de ventilação seguro nos cenários de frio — meta: ZERO.", 1),
        ("Conflitos detectados × conflitos resolvidos sem prejudicar as aves.", 1),
        ("% de tempo de cada eixo dentro da sua zona de conforto.", 1),
        ("Estabilidade (ausência de oscilação liga/desliga dos atuadores).", 1),
        ("Cobertura de testes: nº de casos automatizados que passam.", 1),
    ])

    # 15 — Resultados
    slide_conteudo(prs, "15 · Resultados obtidos", [
        ("Suíte de testes: 22/22 casos passam (física, controladores e arbitragem).", 0, VERDE),
        ("Em TODOS os cenários de frio, o exaustor respeitou o teto seguro da fase — 0 violações.", 0, VERDE),
        ("Conflito frio×amônia: ventilação limitada + aquecedor elevado em compensação, com o conflito explicitado no dashboard.", 0),
        ("Calor + amônia: um único atuador (ventilação) resolveu os dois eixos (sinergia detectada).", 0),
        ("Pintinhos (fase 1) mantidos estáveis em ~32 °C, sem resfriamento indevido.", 0),
        ("Dashboard mostra, em tempo real, o que está certo/errado e o mecanismo de mitigação ativo.", 0, ACCENT),
    ])

    # 16 — Análise crítica
    slide_conteudo(prs, "16 · Análise crítica dos resultados", [
        ("A hierarquia evita matar de frio, MAS aceita NH₃ elevado por mais tempo na fase 1 —", 0, AMAREL),
        ("é um trade-off real, não uma solução mágica: ventilação sozinha não resolve amônia em pintinhos.", 1),
        ("Isso revela um limite físico: na chegada, o controle de amônia depende de MANEJO DE CAMA,", 0),
        ("não de ventilação — algo que o gêmeo deixou evidente.", 1, ACCENT),
        ("O ITU (escala 74/78) NÃO se aplica a pintinhos: a 32 °C o índice já daria 'crítico',", 0, VERMEL),
        ("o que faria o sistema resfriar a própria temperatura que o pinto precisa. Corrigimos isso.", 1),
        ("O modelo físico é simplificado (1ª ordem, constantes calibradas, não validadas com dados reais).", 0),
    ])

    # 17 — Objetivos x resultados
    slide_conteudo(prs, "17 · Comparação: objetivos propostos × alcançados", [
        ("Replicar o firmware em Python — ALCANÇADO (fórmulas e ruído idênticos).", 0, VERDE),
        ("Modelar o ambiente compartilhado com acoplamentos — ALCANÇADO.", 0, VERDE),
        ("Supervisor por bem-estar (térmico>amônia>luz) — ALCANÇADO e testado.", 0, VERDE),
        ("Dashboard de gêmeo digital — ALCANÇADO (certo/errado + mitigação + conflitos).", 0, VERDE),
        ("Validação por cenários e testes — ALCANÇADO (22 testes).", 0, VERDE),
        ("Validação com DADOS REAIS — NÃO alcançado (não era objetivo; é trabalho futuro).", 0, AMAREL),
    ])

    # 18 — Dificuldades
    slide_conteudo(prs, "18 · Principais dificuldades encontradas", [
        ("Acoplamento dos atuadores: perceber que o exaustor é um recurso ÚNICO disputado por dois eixos.", 0),
        ("Conflito aquecedor × exaustor: como ventilar sem congelar — surgiu o conceito de 'teto seguro'.", 0),
        ("ITU em pintinhos: o índice clássico contradiz a curva de aquecimento — exigiu decisão de projeto.", 0),
        ("Calibrar as constantes da física para dinâmica visível e estável na apresentação.", 0),
        ("Rodar 100% offline e sem dependências (substituímos WebSocket por SSE da stdlib).", 0),
    ])

    # 19 — Limitações
    slide_conteudo(prs, "19 · Limitações da solução proposta", [
        ("Modelo físico simplificado e com constantes empíricas (não calibradas com sensores reais).", 0),
        ("Sistema puramente REATIVO: não antecipa, só responde ao estado atual.", 0),
        ("Arbitragem baseada em REGRAS/hierarquia — não é uma otimização multiobjetivo formal.", 0),
        ("Um único galpão/zona homogênea (sem gradientes espaciais de temperatura ou gás).", 0),
        ("Sem persistência de dados nem consideração de custo energético.", 0),
    ])

    # 20 — Pontos positivos
    slide_conteudo(prs, "20 · Pontos positivos do projeto", [
        ("Roda offline, sem internet/broker — confiável para apresentar.", 0, VERDE),
        ("Arquitetura modular (sensores/controladores/supervisor/ambiente) e testada.", 0, VERDE),
        ("O supervisor EXPLICA cada decisão — conflitos e sinergias ficam transparentes.", 0, VERDE),
        ("Fidelidade ao firmware: mesmas fórmulas e ruído do hardware Wokwi.", 0, VERDE),
        ("Dashboard didático: comunica 'o que está certo, o que está errado e o que está sendo feito'.", 0, VERDE),
    ])

    # 21 — Pontos negativos / melhorias
    slide_conteudo(prs, "21 · Pontos negativos e o que poderia melhorar", [
        ("O modelo físico carece de validação com dados reais de um aviário.", 0, AMAREL),
        ("A decisão por regras poderia evoluir para otimização (energia × bem-estar).", 0, AMAREL),
        ("Falta persistência: não há histórico para auditoria ou aprendizado.", 0, AMAREL),
        ("Os ganhos dos controladores foram ajustados empiricamente, não sintonizados formalmente.", 0, AMAREL),
        ("O efeito cruzado luz→calor é estimado; mediria-se melhor com a lâmpada real.", 0, AMAREL),
    ])

    # 22 — Trabalhos futuros
    slide_conteudo(prs, "22 · Trabalhos futuros", [
        ("Base de dados + análise PREDITIVA (séries temporais / ML) — sair do reativo para o proativo.", 0, ACCENT),
        ("Controle preditivo por modelo (MPC) e otimização multiobjetivo (conforto × energia).", 0),
        ("Validação e calibração com sensores reais em um aviário.", 0),
        ("Múltiplas zonas e gradientes espaciais no galpão.", 0),
        ("Integração ao vivo com o hardware Wokwi via MQTT (modo híbrido já previsto).", 0),
        ("Inclusão do manejo de cama como atuador 'lento' contra a amônia.", 0),
    ])

    # 23 — Conclusões
    slide_conteudo(prs, "23 · Conclusões e contribuições", [
        ("Construímos um gêmeo digital reativo coeso para os três eixos da granja.", 0),
        ("Principal contribuição: mostrar que coordenar atuadores que compartilham o ambiente", 0, ACCENT),
        ("exige um SUPERVISOR de bem-estar — três malhas isoladas seriam perigosas.", 1, ACCENT),
        ("O gêmeo tornou explícitos conflitos reais (frio×amônia) e limites físicos (cama × ventilação).", 0),
        ("Contribuição prática: ferramenta didática e base sólida para a evolução preditiva.", 0),
        ("Aprendizado central: em sistemas acoplados, a pergunta certa não é 'como resolvo este eixo?',", 0),
        ("mas 'como esta ação afeta TODO o ambiente e o bem-estar das aves?'.", 1, VERDE),
    ])

    # 24 — Obrigado / referências
    slide_conteudo(prs, "Referências e encerramento", [
        ("Limiares de amônia: Orffa (2023); Precision Poultry Farming / UGA Extension.", 0),
        ("Estresse térmico / ITU: Frontiers in Veterinary Science (2020); índice de Thom.", 0),
        ("Curva de aquecimento: manuais de manejo Cobb / Ross (frango de corte).", 0),
        ("Implementação: pacote gemeo_digital/, dashboard/, hardware_wokwi/, tests/ (22 testes).", 0),
        ("Execução: python run_twin.py  →  http://127.0.0.1:8000  (offline).", 0, ACCENT),
        ("Obrigado!  Perguntas?", 0, VERDE),
    ])


def main():
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    construir(prs)
    saida = os.path.join(os.path.dirname(__file__), "gemeo_digital_granja.pptx")
    prs.save(saida)
    print(f"Slides gerados: {saida}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
